"""
test_blackbox_http.py — Fuzz the two public POST endpoints over HTTP, against
a running copy of the built image rather than the app in-process.

Why this exists: test_fuzz_inputs.py proves the code, but not the container
that ships (uvicorn, the installed dependency versions, the Dockerfile). And
no off-the-shelf scanner reaches the logic behind these endpoints: both take a
JSON document inside a multipart form field, which GitLab DAST and API
Security Testing treat as one opaque string (tasks/security_audit.md, "Dynamic
security testing"). This file sends realistic documents with hostile values
inside them, so the payloads get past json.loads and reach the checks.

Skipped unless BLACKBOX_BASE_URL is set, so the normal test run is unchanged.
In CI the target is the `fuzz:blackbox` job's sidecar, never the live host:

    BLACKBOX_BASE_URL=http://localhost:8000 pytest tests/test_blackbox_http.py

FUZZ_EXAMPLES sets examples per test (default 300). FUZZ_SEED makes a run
random but reproducible; unset, the run is deterministic.

What fails a run:
- any 5xx
- any response slower than REQUEST_TIMEOUT_S (a hang is a finding)
- a response without nosniff, or one that is neither JSON nor, for a
  successful export, a .docx
- an export whose .docx does not open, or whose cell text is not exactly
  the submitted row (escaped, not executed, and not silently dropped)
- the service not answering /api/health after the run
"""

from __future__ import annotations

import io
import json
import os
import re

import httpx
import pytest
from docx import Document
from hypothesis import HealthCheck, given, seed, settings, strategies as st

from tests.fuzz_strategies import (
    RAW_PLACEHOLDER,
    attack_strings,
    mutated,
    png,
    raw_json_tokens,
    seed_manifest,
    seed_metadata,
    with_raw_token,
)

BASE_URL = os.environ.get("BLACKBOX_BASE_URL")
pytestmark = pytest.mark.skipif(not BASE_URL, reason="BLACKBOX_BASE_URL not set")

REQUEST_TIMEOUT_S = 10
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_SEED = os.environ.get("FUZZ_SEED")
_FUZZ = settings(
    max_examples=int(os.environ.get("FUZZ_EXAMPLES", "300")),
    derandomize=not _SEED,
    deadline=None,
    database=None,
    suppress_health_check=[HealthCheck.too_slow, HealthCheck.data_too_large],
)


def fuzz(test):
    test = _FUZZ(test)
    return seed(int(_SEED))(test) if _SEED else test


# ── Preflight: prove the target is reachable and accepts the seeds ───────────
# Without this, a wrong URL or a target that refuses everything would make
# every property below pass: a 4xx is "below 500". That is how API Security
# Testing reported a clean run while F9 was live.

@pytest.fixture(scope="module")
def client():
    with httpx.Client(base_url=BASE_URL, timeout=REQUEST_TIMEOUT_S) as c:
        health = c.get("/api/health")
        assert health.status_code == 200, f"target not healthy: {health.status_code}"
        r = _evaluate(c, json.dumps(seed_metadata()))
        assert r.status_code == 200, f"seed metadata refused: {r.status_code} {r.text[:300]}"
        r = _export(c, json.dumps(seed_manifest()))
        assert r.status_code == 200, f"seed manifest refused: {r.status_code} {r.text[:300]}"
        yield c


def _evaluate(c: httpx.Client, body: str) -> httpx.Response:
    return c.post("/api/evaluate", data={"metadata_json": body})


def _export(c: httpx.Client, body: str) -> httpx.Response:
    return c.post("/api/export/docx", data={"manifest_json": body},
                  files=[("crops", ("c0.png", png(), "image/png"))])


# ── Pass/fail rules shared by every property ─────────────────────────────────

def _assert_safe(r: httpx.Response, body: str) -> None:
    assert r.status_code < 500, (r.status_code, body[:500])
    assert r.headers.get("x-content-type-options") == "nosniff", (r.status_code, dict(r.headers))
    content_type = r.headers.get("content-type", "")
    if r.status_code == 200 and content_type.startswith(DOCX_TYPE):
        Document(io.BytesIO(r.content))  # raises if the file is not a valid .docx
    else:
        assert content_type.startswith("application/json"), (r.status_code, content_type)


# XML 1.0's legal characters (spec production [2]). The export must drop
# exactly the rest: anything else in a row's text has to come back verbatim.
_XML_ILLEGAL = re.compile("[^\t\n\r\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")


def _assert_rows_round_trip(r: httpx.Response, body: str) -> None:
    # Expected text comes from what the server decodes, not from the Python
    # object we serialised: two generated lone surrogates "\ud800" "\udc00"
    # are, once escaped into JSON, one valid pair that decodes to U+10000.
    manifest = json.loads(body)
    if r.status_code != 200 or not isinstance(manifest, list):
        return
    rows = [row for row in manifest if isinstance(row, dict)]
    if not rows:
        return  # the export writes a placeholder row instead
    fields = ("status", "check_title", "check_id", "text")
    if not all(isinstance(row.get(f, ""), str) for row in rows for f in fields):
        return  # type mutations: the exact text is not well-defined
    table = Document(io.BytesIO(r.content)).tables[0]
    for row, table_row in zip(rows, table.rows[1:], strict=True):
        expected = f"[{row.get('status', '')}] {row.get('check_title', '')} " \
                   f"({row.get('check_id', '')}): {row.get('text', '')}"
        # python-docx writes \r as a line break, which reads back as \n.
        expected = _XML_ILLEGAL.sub("", expected).replace("\r", "\n")
        assert table_row.cells[1].text == expected, (row, table_row.cells[1].text[:300])


def _is_str(node) -> bool:
    return isinstance(node, str)


# ── 1. Wrong types (the in-process strategy, now through the real image) ─────

@fuzz
@given(mutated(seed_metadata))
def test_evaluate_type_mutations(client, metadata):
    body = json.dumps(metadata)
    _assert_safe(_evaluate(client, body), body)


@fuzz
@given(mutated(seed_manifest))
def test_export_type_mutations(client, manifest):
    body = json.dumps(manifest)
    _assert_safe(_export(client, body), body)


# ── 2. Hostile text in string fields ─────────────────────────────────────────

@fuzz
@given(mutated(seed_metadata, values=attack_strings, where=_is_str))
def test_evaluate_attack_strings(client, metadata):
    body = json.dumps(metadata)
    _assert_safe(_evaluate(client, body), body)


@fuzz
@given(mutated(seed_manifest, values=attack_strings, where=_is_str))
def test_export_attack_strings(client, manifest):
    body = json.dumps(manifest)
    r = _export(client, body)
    _assert_safe(r, body)
    _assert_rows_round_trip(r, body)


# ── 3. Parser extremes: huge numbers, deep nesting, NaN, duplicate keys ──────

@fuzz
@given(mutated(seed_metadata, values=st.just(RAW_PLACEHOLDER)), raw_json_tokens)
def test_evaluate_parser_extremes(client, metadata, token):
    body = with_raw_token(metadata, token)
    _assert_safe(_evaluate(client, body), body)


@fuzz
@given(mutated(seed_manifest, values=st.just(RAW_PLACEHOLDER)), raw_json_tokens)
def test_export_parser_extremes(client, manifest, token):
    body = with_raw_token(manifest, token)
    _assert_safe(_export(client, body), body)


# ── Last: the service survived everything above ──────────────────────────────

def test_zz_service_still_healthy(client):
    assert client.get("/api/health").status_code == 200
